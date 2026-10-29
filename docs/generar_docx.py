"""Convierte docs/MANUAL_TECNICO.md en un documento Word (docs/Manual_tecnico.docx).

    .venv\\Scripts\\python docs\\generar_docx.py

Soporta el subconjunto de Markdown usado en el manual: títulos, párrafos, listas (con un nivel de
anidamiento), tablas, bloques de código, citas, negrita y código en línea.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
ORIGEN = AQUI / "MANUAL_TECNICO.md"
DESTINO = AQUI / "Manual_tecnico.docx"
EN_LINEA = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def sombrear(elemento, color):
    pr = elemento.get_or_add_tcPr() if hasattr(elemento, "get_or_add_tcPr") else elemento.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pr.append(shd)


def texto_en_linea(parrafo, texto, tam=None):
    for parte in EN_LINEA.split(texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = parrafo.add_run(parte[2:-2].replace("`", ""))
            r.bold = True
        elif parte.startswith("`"):
            r = parrafo.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
        else:
            r = parrafo.add_run(parte)
        if tam:
            r.font.size = Pt(tam)


def tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    cabecera, cuerpo = celdas[0], [f for f in celdas[2:]]
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        p = t.rows[0].cells[i].paragraphs[0]
        texto_en_linea(p, f"**{c}**" if c else "", 9)
        sombrear(t.rows[0].cells[i]._tc, "E4E3DF")
    for fila in cuerpo:
        cel = t.add_row().cells
        for i in range(len(cabecera)):
            texto_en_linea(cel[i].paragraphs[0], fila[i] if i < len(fila) else "", 9)
    doc.add_paragraph()


def codigo(doc, lineas):
    p = doc.add_paragraph()
    sombrear(p._p, "F0EFEC")
    p.paragraph_format.left_indent = Cm(0.3)
    for i, linea in enumerate(lineas):
        r = p.add_run(linea)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        if i < len(lineas) - 1:
            r.add_break()


def main():
    lineas = ORIGEN.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    pie = doc.sections[0].footer.paragraphs[0]
    pie.text = "Sistema inteligente de monitoreo satelital y alerta temprana de contaminación y amenazas forestales · Huánuco · Manual técnico"
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.runs[0].font.size = Pt(8)

    primer_titulo = True
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        if linea.startswith("```"):
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].startswith("```"):
                bloque.append(lineas[i])
                i += 1
            codigo(doc, bloque)
        elif linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(lineas[i])
                i += 1
            tabla(doc, filas)
            continue
        elif m := re.match(r"^(#{1,4}) (.+)", linea):
            nivel, texto = len(m.group(1)), m.group(2).replace("`", "")
            if nivel == 1 and primer_titulo:
                doc.add_heading(texto, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
                primer_titulo = False
            else:
                if nivel == 1:  # cada parte empieza en página nueva
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                doc.add_heading(texto, min(nivel, 3))
        elif linea.strip() == "---":
            pass
        elif linea.startswith("> "):
            p = doc.add_paragraph()
            sombrear(p._p, "FDF3D8")
            p.paragraph_format.left_indent = Cm(0.4)
            texto_en_linea(p, linea[2:])
        elif m := re.match(r"^( *)- (.+)", linea):
            estilo = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            texto_en_linea(doc.add_paragraph(style=estilo), m.group(2))
        elif m := re.match(r"^( *)(\d+)\. (.+)", linea):
            # Número literal: el estilo «List Number» de Word continuaría la numeración entre listas
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.4 if len(m.group(1)) >= 2 else 0.7)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            texto_en_linea(p, f"{m.group(2)}. {m.group(3)}")
        elif linea.strip():
            texto_en_linea(doc.add_paragraph(), linea)
        i += 1

    doc.save(DESTINO)
    print(f"Manual: {DESTINO}")


if __name__ == "__main__":
    main()
