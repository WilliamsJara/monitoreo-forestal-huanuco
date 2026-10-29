"""Glosario detallado del proyecto en Word (docs/Glosario.docx).

    .venv\\Scripts\\python docs\\generar_glosario.py

Cada término tiene su definición general y cómo se usa en este sistema, agrupados por tema.
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESTINO = Path(__file__).resolve().parent / "Glosario.docx"
TITULO = ("Sistema inteligente de monitoreo satelital y alerta temprana de contaminación y amenazas forestales "
          "en el departamento de Huánuco")

# (tema, [(término, definición, uso en el proyecto)])
GLOSARIO = [
    ("1. Amenazas ambientales y cobertura forestal", [
        ("Alerta temprana", "Aviso que se emite apenas aparecen indicios de un daño, antes de que se agrave, para "
         "que las autoridades puedan actuar a tiempo.",
         "Objetivo central del sistema: avisos por provincia sobre incendios, deforestación, vegetación y aire."),
        ("Amenaza forestal", "Proceso que destruye o degrada el bosque: tala, quema, desbroce, incendios.",
         "El sistema vigila incendios y deforestación como amenazas forestales."),
        ("Deforestación", "Conversión del bosque en otro uso de la tierra (pastos, cultivos, caminos) por "
         "acción humana.", "Se detecta como pérdida de bosque dentro del mapa oficial de MINAM."),
        ("Tala", "Corte de árboles, selectivo o total.", "Aparece como claro nuevo en las imágenes Sentinel-2."),
        ("Quema", "Uso del fuego para eliminar vegetación, a menudo tras la tala.",
         "Se ve como foco de calor (FIRMS) y como cicatriz oscura en falso color."),
        ("Desbroce", "Eliminación de la vegetación baja o del monte para preparar el terreno.",
         "Una de las causas de pérdida que agrupa la alerta de deforestación."),
        ("Pérdida de cobertura forestal", "Disminución de la superficie con bosque en un periodo.",
         "Píxeles de bosque cuyo NDVI cae ≥ 0,25 y queda ≤ 0,55 frente al mismo periodo del año anterior."),
        ("Recuperación (regeneración)", "Aumento de la vegetación en un área antes degradada.",
         "Píxeles que alguna vez fueron bosque, con NDVI ≤ 0,45 que sube ≥ 0,15."),
        ("Bosque húmedo amazónico", "Bosque tropical de la Amazonía, de selva baja y selva alta.",
         "Capa «bosque al 2024» de MINAM, usada como máscara de bosque."),
        ("Bosque estacionalmente seco", "Bosque que pierde hojas en la época seca (valles interandinos, como el "
         "del Marañón).", "Capa «bosque seco 2021» de MINAM, sumada a la máscara de bosque."),
        ("Bosque primario", "Bosque que no ha sido talado ni alterado de forma importante.",
         "Es el bosque que mapea MINAM; la pérdida se evalúa sobre él."),
        ("Bosque secundario (purma)", "Vegetación que vuelve a crecer en un área ya deforestada.",
         "No está en el mapa oficial: su tala no se evalúa (limitación documentada)."),
        ("Cambio agrícola", "Variación por cosecha, rotación o preparación de parcelas, que no es pérdida de "
         "bosque.", "Causa de falsas alarmas eliminada con la máscara de bosque de MINAM."),
        ("Estacionalidad", "Variación natural a lo largo del año (época seca y de lluvias).",
         "Se evita compararla con daño: se compara siempre con el mismo periodo del año anterior."),
        ("Contaminación atmosférica", "Presencia en el aire de gases o partículas dañinas.",
         "Se vigila con NO₂ y CO de Sentinel-5P; es parte del título del proyecto."),
    ]),
    ("2. Satélites, sensores y datos", [
        ("Teledetección", "Obtención de información de la superficie terrestre a distancia, desde satélites o "
         "aviones.", "Base de todo el sistema."),
        ("Sentinel-2", "Par de satélites europeos (programa Copernicus) con cámara multiespectral de 10–60 m, "
         "que pasan cada ~5 días.", "Fuente para deforestación, recuperación, NDVI y las imágenes de evidencia."),
        ("Sentinel-2 L2A", "Producto de Sentinel-2 corregido atmosféricamente (reflectancia de superficie), con "
         "clasificación de escena (SCL).", "Colección «sentinel-2-l2a» que consulta el sistema."),
        ("Sentinel-5P / TROPOMI", "Satélite europeo con el espectrómetro TROPOMI, que mide gases atmosféricos a "
         "~5 km.", "Fuente de NO₂ y CO semanales por provincia."),
        ("Sentinel-1", "Satélites europeos de radar, que ven a través de las nubes.",
         "Próximo paso propuesto para la selva alta nubosa."),
        ("Landsat", "Satélites de la NASA/USGS con imágenes de 30 m desde 1972.",
         "Base de las capas de MINAM: por eso la validación es independiente de Sentinel-2."),
        ("VIIRS", "Sensor de los satélites S-NPP, NOAA-20 y NOAA-21 que detecta fuegos a 375 m.",
         "Principal fuente de focos de calor (tres satélites)."),
        ("MODIS", "Sensor de los satélites Terra y Aqua que detecta fuegos a 1 km.",
         "Cuarta fuente de focos de calor."),
        ("NRT (casi tiempo real)", "Datos publicados pocas horas después de la toma.",
         "Los focos de FIRMS son NRT: permiten alertas del mismo día."),
        ("Resolución espacial", "Tamaño en el terreno que representa un píxel.",
         "Entre ~33 m (Yarowilca) y ~115 m (Puerto Inca) en el análisis de cambio."),
        ("Banda espectral", "Rango de longitudes de onda que registra el sensor (azul, rojo, infrarrojo…).",
         "B02 azul, B03 verde, B04 rojo, B08 infrarrojo cercano, B8A, B12 infrarrojo de onda corta."),
        ("Reflectancia", "Proporción de luz que refleja la superficie (0 a 1).",
         "Valor de cada banda con el que se calcula el NDVI."),
        ("NIR (infrarrojo cercano)", "Luz invisible que la vegetación sana refleja mucho.", "Banda B08 del NDVI."),
        ("SWIR (infrarrojo de onda corta)", "Luz que distingue suelo desnudo, quemas y humedad.",
         "Banda B12 de las imágenes en falso color."),
        ("Falso color", "Imagen que asigna bandas invisibles a colores visibles.",
         "B12-B8A-B04: bosque verde oscuro; suelo desnudo o quema en rosado o marrón."),
        ("Color real", "Imagen con las bandas roja, verde y azul, como la vería el ojo.",
         "Imágenes de antes y después en las láminas y en el tablero."),
        ("Escena", "Una toma del satélite en una fecha.", "Las láminas antiguas usaban una sola escena (con nubes)."),
    ]),
    ("3. Indicadores y variables", [
        ("NDVI", "Índice de vegetación de diferencia normalizada: (NIR − rojo) / (NIR + rojo). Va de −1 a 1; "
         "más alto = vegetación más densa.", "Indicador de bosque (≥ 0,60), de pérdida y de salud de la vegetación."),
        ("Foco de calor", "Píxel donde el satélite detecta una temperatura anormalmente alta, normalmente fuego.",
         "Unidad básica de los incendios (NASA FIRMS)."),
        ("FRP (potencia radiativa del fuego)", "Energía que emite el fuego, en megavatios (MW).",
         "Mide la intensidad; define la severidad de eventos y alertas."),
        ("Confianza (FIRMS)", "Probabilidad de que el foco sea fuego real: baja, nominal o alta.",
         "Se informa como calidad de la alerta de incendio."),
        ("NO₂ (dióxido de nitrógeno)", "Gas de la combustión (vehículos, industria, quemas).",
         "Columna troposférica semanal por provincia, en µmol/m²."),
        ("CO (monóxido de carbono)", "Gas de la combustión incompleta; sube con quemas e incendios.",
         "Columna total semanal por provincia, en mmol/m²."),
        ("Columna atmosférica", "Cantidad total de un gas en la columna de aire sobre un punto (no la "
         "concentración al nivel del suelo).", "Lo que mide Sentinel-5P; limitación documentada."),
        ("Hectárea (ha)", "Superficie de 10 000 m² (100 × 100 m).", "Unidad de la pérdida de bosque."),
        ("Área observable", "Porcentaje de la provincia sin nubes en ambos periodos comparados.",
         "Calidad de la alerta de deforestación; aviso si es menor al 40 %."),
    ]),
    ("4. Procesamiento de imágenes", [
        ("Píxel", "Unidad mínima de una imagen: un cuadrado del terreno con un valor por banda.",
         "El cambio se calcula píxel a píxel."),
        ("Máscara", "Capa que indica qué píxeles se usan y cuáles se descartan.",
         "Máscara de nubes y máscara de bosque."),
        ("SCL", "Clasificación de escena de Sentinel-2 (vegetación, suelo, agua, nube, sombra…).",
         "Solo se aceptan las clases 4 (vegetación) y 5 (suelo)."),
        ("s2cloudless", "Modelo de aprendizaje automático que estima la probabilidad de nube en Sentinel-2.",
         "Eliminó las 9 de 9 falsas alarmas por nubes de la muestra."),
        ("CLM / CLP", "Bandas de s2cloudless: CLM = máscara de nube (incluye margen alrededor); CLP = "
         "probabilidad de nube (0–255).", "Observación válida si CLM = 0 y CLP < 51 (20 %)."),
        ("Observación válida", "Pasada del satélite en que el píxel se vio sin nubes, sombras, agua ni nieve.",
         "Se exigen al menos 2 por periodo y por píxel."),
        ("Mosaico ORBIT", "Modo de Sentinel Hub que entrega todas las pasadas del periodo para cada píxel.",
         "Permite calcular la mediana de todas las observaciones sin nubes."),
        ("Mosaico simple / leastCC", "Modo que toma una sola escena por píxel (la más reciente o la menos "
         "nublada).", "Se reemplazó por ORBIT: en selva casi siempre devolvía nubes."),
        ("Mediana multitemporal", "Valor central de las observaciones de varias fechas; resiste valores "
         "extremos.", "NDVI de cada periodo y composiciones sin nubes."),
        ("Composición sin nubes", "Imagen armada con la mediana de las pasadas despejadas.",
         "Imágenes de evidencia de antes y después."),
        ("Análisis multitemporal", "Comparación de la misma zona en distintas fechas.",
         "Periodo reciente de 90 días frente al mismo periodo del año anterior."),
        ("Píxel aislado", "Píxel marcado sin vecinos marcados; suele ser ruido.",
         "Se descarta si tiene menos de 3 vecinos marcados (ventana 3 × 3)."),
        ("Zona de pérdida (zona prioritaria)", "Grupo de píxeles de pérdida dentro de una celda de 16 × 16 "
         "píxeles, con al menos 1 ha.", "Unidad de las tablas, de la validación y de las verificaciones de campo."),
        ("Evalscript", "Pequeño programa en JavaScript que Sentinel Hub ejecuta sobre cada píxel.",
         "Calcula NDVI, máscaras y medianas en el servidor de Copernicus."),
    ]),
    ("5. Inteligencia artificial, estadística y validación", [
        ("Inteligencia artificial (IA)", "Técnicas que permiten a un sistema realizar tareas que requieren "
         "criterio: clasificar, agrupar, detectar patrones.", "Aprendizaje automático (nubes), agrupamiento "
         "(incendios), detección de anomalías (aire) y calibración con datos."),
        ("Aprendizaje automático", "Rama de la IA en que un modelo aprende patrones a partir de datos.",
         "s2cloudless clasifica nubes en cada imagen."),
        ("Aprendizaje profundo", "Aprendizaje automático con redes neuronales de muchas capas.",
         "Próximo paso, cuando existan datos verificados para entrenarlo."),
        ("Agrupamiento (clustering) no supervisado", "Formar grupos de datos parecidos sin ejemplos "
         "etiquetados.", "DBSCAN agrupa focos de calor en eventos."),
        ("DBSCAN", "Algoritmo de agrupamiento por densidad: une puntos con suficientes vecinos cercanos.",
         "Focos a menos de 1,5 km, con al menos 3 focos, forman un evento de incendio."),
        ("Distancia haversine", "Distancia sobre la superficie de la esfera terrestre entre dos coordenadas.",
         "Métrica de DBSCAN para los focos."),
        ("Evento de incendio", "Grupo de focos cercanos que corresponden a un mismo fuego.",
         "Tiene severidad alta, media o baja según focos, FRP y días activo."),
        ("Anomalía", "Valor que se aparta de lo habitual.", "Semana con NO₂ o CO muy por encima de su línea base."),
        ("Puntaje z", "Cuántas desviaciones estándar se aleja un valor del promedio.",
         "z ≥ 2 = alerta media; z ≥ 3 = alerta alta de aire."),
        ("Línea base", "Comportamiento normal de referencia.", "Semanas anteriores de la misma provincia."),
        ("Umbral", "Valor límite a partir del cual se toma una decisión.",
         "NDVI de bosque 0,60; caída 0,25; NDVI tras la pérdida 0,55; etc. (config.py)."),
        ("Calibración", "Ajuste de umbrales para que el sistema acierte mejor.",
         "El umbral 0,55 se eligió midiendo precisión y sensibilidad frente a MINAM."),
        ("Precisión", "De lo que el sistema marca, qué parte es correcta.",
         "95 % de las zonas de pérdida coincide con alertas MINAM."),
        ("Sensibilidad (exhaustividad)", "De lo que realmente ocurrió, qué parte detecta el sistema.",
         "28 % del área de alertas MINAM observable (40 % en claros grandes)."),
        ("Falso positivo (falsa alarma)", "Alerta que no corresponde a un hecho real.",
         "Se redujeron de 13 de 20 a 0 en la muestra revisada."),
        ("Falso negativo (omisión)", "Hecho real que el sistema no detecta.", "Medido con la sensibilidad."),
        ("Coincidencia por azar", "Acierto que tendría un sistema que marcara zonas al azar.",
         "11 %: sirve para demostrar que el 95 % no es casualidad."),
        ("Validación independiente", "Comprobación con una fuente distinta de la que usa el sistema.",
         "Comparación con alertas de MINAM (otro satélite y otro método)."),
        ("Tolerancia (±1 píxel)", "Margen de posición aceptado al comparar dos mapas.",
         "Compensa diferencias de georreferenciación entre Landsat y Sentinel-2."),
        ("Muestra aleatoria y semilla", "Selección al azar de casos; la semilla permite repetir el mismo sorteo.",
         "Muestras de 20, 30 y 30 zonas (semillas 42, 2026 y 2027)."),
        ("Fotointerpretación", "Interpretación visual de imágenes satelitales por una persona.",
         "Revisión de las láminas; la actual fue asistida por IA y es preliminar."),
        ("Verdad de campo", "Verificación en el terreno.", "Próximo paso con autoridades forestales."),
        ("Intervalo de confianza (IC 95 %)", "Rango donde está el valor real con 95 % de probabilidad.",
         "Calculado con el método de Wilson para las muestras."),
        ("Sobreestimación optimista", "Resultado mejor de lo real por elegir parámetros mirando la misma "
         "referencia con la que se evalúa.", "Por eso el umbral 0,55 se verificó también con una muestra visual."),
    ]),
    ("6. Fuentes oficiales y plataformas", [
        ("Copernicus", "Programa de observación de la Tierra de la Unión Europea.", "Provee Sentinel-2 y Sentinel-5P."),
        ("Copernicus Data Space Ecosystem (CDSE)", "Plataforma gratuita de acceso y procesamiento de datos "
         "Copernicus.", "Servicio que consulta el sistema con credenciales OAuth."),
        ("Sentinel Hub", "Motor de procesamiento de imágenes dentro de CDSE.", "Lo usan las API de proceso y estadística."),
        ("Process API", "Servicio que devuelve imágenes procesadas.", "Imagen de cambio, color real y recortes."),
        ("Statistical API", "Servicio que devuelve estadísticas por área y periodo.", "NDVI mensual y aire semanal."),
        ("Unidades de procesamiento (PU)", "Medida del consumo de cómputo en Sentinel Hub.",
         "La cuenta gratuita tiene cuota mensual; se renueva el día 1 de cada mes."),
        ("Cuota", "Límite de uso asignado a cada cuenta.",
         "Se agotó el 24/10/2026; error 403 «ACCESS_INSUFFICIENT_PROCESSING_UNITS»."),
        ("OAuth / Client ID y Client secret", "Protocolo de autorización; el par identifica y autentica a la "
         "aplicación.", "Credenciales de Copernicus en el archivo .env."),
        ("NASA FIRMS", "Sistema de la NASA que distribuye focos de calor de VIIRS y MODIS.",
         "Fuente de incendios; se consulta con una MAP_KEY."),
        ("MAP_KEY", "Clave de acceso gratuita a la API de FIRMS.", "Variable FIRMS_MAP_KEY del .env."),
        ("MINAM", "Ministerio del Ambiente del Perú.", "Fuente oficial de bosque y alertas."),
        ("PNCB", "Programa Nacional de Conservación de Bosques para la Mitigación del Cambio Climático (MINAM).",
         "Produce las alertas tempranas oficiales de deforestación."),
        ("Geobosques", "Plataforma de MINAM/PNCB con información de bosque, pérdida y alertas.",
         "Servicios abiertos usados como máscara y como validación."),
        ("Alerta temprana MINAM", "Aviso oficial de pérdida reciente de bosque, basado en Landsat.",
         "Referencia independiente: columna «alerta MINAM: sí/no» del tablero."),
        ("SERFOR / ATFFS", "Servicio Nacional Forestal y de Fauna Silvestre / sus administraciones técnicas "
         "regionales.", "Posibles usuarios del sistema."),
        ("geoBoundaries", "Base abierta de límites administrativos del mundo.",
         "Origen de los límites de las 11 provincias (nivel ADM2)."),
    ]),
    ("7. Geografía y datos espaciales", [
        ("Coordenadas (latitud, longitud)", "Posición de un punto en la Tierra, en grados.",
         "Cada alerta y zona trae sus coordenadas."),
        ("WGS84 / CRS84", "Sistema de referencia geográfico estándar (el del GPS).",
         "Sistema de todas las coordenadas del proyecto."),
        ("Bbox (caja envolvente)", "Rectángulo mínimo que contiene una geometría.",
         "Área que se pide a los servicios para cada provincia."),
        ("GeoJSON", "Formato de texto para geometrías geográficas.", "data/provincias_huanuco.geojson."),
        ("ADM2", "Segundo nivel administrativo (en Perú, la provincia).", "Nivel de las 11 provincias de Huánuco."),
        ("Ubigeo", "Código oficial de ubicación geográfica del Perú (departamento, provincia, distrito).",
         "Campo de las alertas de MINAM (Huánuco = 10)."),
        ("Capa", "Conjunto de datos geográficos que se superpone en un mapa.",
         "Provincias, focos, eventos, bosque y alertas MINAM en el tablero."),
        ("WMS", "Estándar para publicar mapas como imágenes por internet.",
         "Capas de MINAM que se pueden activar en el mapa del tablero."),
    ]),
    ("8. Alertas del sistema", [
        ("Nivel de alerta", "Gravedad: alta, media, baja o sin alerta.",
         "Cada provincia toma el nivel más alto de sus alertas."),
        ("Evidencia de la alerta", "Datos que permiten verificar una alerta.",
         "Fuente, fecha, indicador que la activó, calidad de la observación y coordenadas."),
        ("Indicador que activa la alerta", "Medida y umbral que dispararon el aviso.",
         "Ej.: «n.º de focos ≥ 25 = alta»; «caída de NDVI ≥ 0,25 en bosque»."),
        ("Calidad de la observación", "Qué tan confiable fue la medición.",
         "Ej.: % de la provincia sin nubes, % de focos con confianza alta."),
        ("Severidad", "Gravedad de un evento de incendio.", "Alta si ≥ 15 focos o ≥ 300 MW."),
        ("Corrida", "Ejecución completa del sistema que genera datos y alertas.",
         "Se programa a diario; cada corrida se guarda en la base de datos."),
    ]),
    ("9. Software e infraestructura", [
        ("Python", "Lenguaje de programación.", "Lenguaje de todo el sistema."),
        ("Entorno virtual (.venv)", "Carpeta con Python y librerías aisladas para un proyecto.",
         "Se crea en cada equipo; no se copia entre computadoras."),
        ("Streamlit", "Librería de Python para crear aplicaciones web de datos.", "Tecnología del tablero."),
        ("Tablero (dashboard)", "Pantalla que resume indicadores, mapas y alertas.",
         "Interfaz web del sistema, adaptada a celular."),
        ("Modo guardado / modo en vivo", "Mostrar la última corrida guardada o consultar los satélites al momento.",
         "El tablero abre en modo guardado (segundos, sin gastar cuota)."),
        ("Caché", "Copia guardada de un resultado para no recalcularlo.",
         "Carpeta .cache; NDVI y cobertura se recalculan una vez por semana."),
        ("TTL (tiempo de vida)", "Tiempo que un dato en caché se considera vigente.", "3 h focos; 24 h aire; 7 días NDVI y cobertura."),
        ("Fuentes independientes", "Cada fuente se consulta y actualiza por separado.",
         "Si una falla, se usan sus últimos datos buenos (respaldo por fuente) con su fecha."),
        ("Base de datos SQLite", "Base de datos en un solo archivo, sin servidor.",
         "salidas/monitoreo.sqlite: corridas, alertas con evidencia y zonas."),
        ("Respaldo (copia de seguridad)", "Copia de los datos para recuperarlos si se pierden.",
         "ZIP diario verificado: se comprueba que se puede restaurar."),
        ("API", "Interfaz para que un programa pida datos o servicios a otro.",
         "Copernicus, FIRMS y Geobosques se consultan por API."),
        ("CSV", "Archivo de tabla en texto separado por comas.", "Resultados exportables de cada corrida."),
        ("Programador de tareas / cron", "Herramienta de Windows / Linux para ejecutar tareas a una hora fija.",
         "Actualización diaria a las 06:00 y respaldo a las 06:30."),
        ("Servidor Ubuntu", "Computadora con Linux Ubuntu que publica el sistema en internet.",
         "Destino de la migración: http://149.34.48.35/carpeta/."),
        ("systemd (servicio)", "Gestor de servicios de Linux: arranca programas y los reinicia si fallan.",
         "Servicio «monitoreo-huanuco» que mantiene el tablero activo."),
        ("nginx (proxy inverso)", "Servidor web que recibe las visitas y las reenvía al tablero.",
         "Publica el tablero en la ruta /carpeta."),
        ("Variables de entorno (.env)", "Archivo con configuración y claves fuera del código.",
         "Credenciales de Copernicus y FIRMS; no se comparte."),
    ]),
]


def sombrear(celda, color):
    pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    pr.append(shd)


def repetir_encabezado(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    tr_pr.append(el)


def main():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"; estilo.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
        pie = s.footer.paragraphs[0]
        pie.text = "Glosario · Sistema de monitoreo satelital y alerta temprana — Huánuco · INNOBEST HUÁNUCO 2026"
        pie.runs[0].font.size = Pt(8); pie.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    t = doc.add_heading("Glosario de términos", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(TITULO); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p = doc.add_paragraph("Williams Roberts Jara Daza · Limber Yinier Valdivia Cuzque\n"
                          f"Ingeniería de software con inteligencia artificial · {date.today():%d/%m/%Y}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = sum(len(t) for _, t in GLOSARIO)
    doc.add_paragraph(f"Este glosario reúne los {total} términos técnicos usados en el sistema, sus documentos y "
                      "su presentación. Para cada término se da su significado general y cómo se aplica en este "
                      "proyecto. Los términos están agrupados por tema.")

    for tema, terminos in GLOSARIO:
        doc.add_heading(tema, 1)
        tabla = doc.add_table(rows=1, cols=3)
        tabla.style = "Table Grid"
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, texto in zip(tabla.rows[0].cells, ["Término", "Definición", "En este proyecto"]):
            c.text = ""
            r = c.paragraphs[0].add_run(texto); r.bold = True; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            sombrear(c, "0F3D2E")
        repetir_encabezado(tabla.rows[0])
        for i, (termino, definicion, uso) in enumerate(sorted(terminos, key=lambda x: x[0].lower())):
            celdas = tabla.add_row().cells
            for c, texto, negrita in zip(celdas, (termino, definicion, uso), (True, False, False)):
                c.text = ""
                r = c.paragraphs[0].add_run(texto); r.bold = negrita; r.font.size = Pt(9.5)
                if i % 2:
                    sombrear(c, "F2F5F3")
        for fila in tabla.rows:
            for c, w in zip(fila.cells, (Cm(3.6), Cm(6.9), Cm(6.5))):
                c.width = w

    doc.save(DESTINO)
    print(f"{DESTINO} ({total} términos)")


if __name__ == "__main__":
    main()
