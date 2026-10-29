"""Presentación para la exposición en INNOBEST HUÁNUCO 2026 (PowerPoint 16:9 con notas del orador).

    .venv\\Scripts\\python scripts\\presentacion.py

Usa los resultados ya calculados (validación MINAM, última corrida, figuras de material_innobest.py y
capturas del tablero en salidas/innobest/capturas) y genera salidas/innobest/Presentacion_INNOBEST_2026.pptx.
"""
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from huanuco_monitor import almacen, config  # noqa: E402

TITULO = ("Sistema inteligente de monitoreo satelital y alerta temprana de contaminación y amenazas forestales "
          "en el departamento de Huánuco")
EQUIPO = [("Williams Roberts Jara Daza", "Ingeniería de software con inteligencia artificial"),
          ("Limber Yinier Valdivia Cuzque", "Ingeniería de software con inteligencia artificial")]
INN = config.SALIDAS_DIR / "innobest"
CAP = INN / "capturas"
VAL = config.SALIDAS_DIR / "validacion"
TMP = INN / "recortes"

VERDE_OSC = RGBColor(0x0F, 0x3D, 0x2E)
VERDE = RGBColor(0x1F, 0x7A, 0x4D)
NARANJA = RGBColor(0xEB, 0x68, 0x34)
AZUL = RGBColor(0x2A, 0x78, 0xD6)
TINTA = RGBColor(0x1B, 0x1B, 0x1A)
TINTA_2 = RGBColor(0x52, 0x51, 0x4E)
APAGADO = RGBColor(0x89, 0x87, 0x81)
FONDO = RGBColor(0xFA, 0xFA, 0xF8)
LINEA = RGBColor(0xE1, 0xE0, 0xD9)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]
ANCHO, ALTO = prs.slide_width, prs.slide_height


# --------------------------------------------------------------------------- utilidades

def texto(slide, x, y, w, h, contenido, tam=16, color=TINTA, negrita=False, alinear=PP_ALIGN.LEFT,
          anclaje=MSO_ANCHOR.TOP, interlineado=1.1):
    """`contenido`: str o lista de párrafos; cada párrafo puede ser str o lista de (texto, dict de formato)."""
    caja = slide.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anclaje
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        p.line_spacing = interlineado
        trozos = par if isinstance(par, list) else [(par, {})]
        for t, fmt in trozos:
            r = p.add_run()
            r.text = t
            r.font.name = FUENTE
            r.font.size = Pt(fmt.get("tam", tam))
            r.font.bold = fmt.get("negrita", negrita)
            r.font.color.rgb = fmt.get("color", color)
    return caja


def rect(slide, x, y, w, h, relleno, borde=None, forma=MSO_SHAPE.RECTANGLE, radio=None):
    s = slide.shapes.add_shape(forma, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = relleno
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radio is not None and forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    return s


def imagen(slide, ruta, x, y, w=None, h=None, borde=True):
    pic = slide.shapes.add_picture(str(ruta), x, y, w, h)
    if borde:
        pic.line.color.rgb = LINEA
        pic.line.width = Pt(0.75)
    return pic


def imagen_ajustada(slide, ruta, x, y, w_max, h_max, borde=True):
    """Coloca la imagen dentro de la caja manteniendo su proporción, centrada."""
    ancho, alto = Image.open(ruta).size
    escala = min(w_max / ancho, h_max / alto)
    w, h = int(ancho * escala), int(alto * escala)
    return imagen(slide, ruta, x + (w_max - w) // 2, y, w, h, borde)


def recortar(origen, caja, nombre):
    TMP.mkdir(parents=True, exist_ok=True)
    destino = TMP / nombre
    Image.open(origen).crop(caja).save(destino)
    return destino


numero = {"n": 0}


def diapositiva(titulo, seccion, notas):
    numero["n"] += 1
    s = prs.slides.add_slide(VACIA)
    rect(s, 0, 0, ANCHO, ALTO, FONDO)
    rect(s, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.75), NARANJA)
    texto(s, Inches(0.85), Inches(0.5), Inches(11), Inches(0.3), seccion.upper(), tam=11, color=NARANJA,
          negrita=True)
    texto(s, Inches(0.85), Inches(0.78), Inches(11.8), Inches(0.7), titulo, tam=28, color=VERDE_OSC, negrita=True)
    rect(s, Inches(0.6), Inches(7.0), Inches(12.13), Emu(9525), LINEA)
    texto(s, Inches(0.6), Inches(7.08), Inches(10), Inches(0.3),
          "INNOBEST HUÁNUCO 2026  ·  Monitoreo satelital y alerta temprana — Huánuco", tam=10, color=APAGADO)
    texto(s, Inches(11.73), Inches(7.08), Inches(1), Inches(0.3), str(numero["n"]), tam=10, color=APAGADO,
          alinear=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notas
    return s


def tarjeta_dato(s, x, y, w, h, cifra, etiqueta, color=VERDE_OSC, tam_cifra=40):
    rect(s, x, y, w, h, BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
    texto(s, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.9), cifra, tam=tam_cifra, color=color,
          negrita=True)
    texto(s, x + Inches(0.25), y + Inches(0.2) + Pt(tam_cifra * 1.35), w - Inches(0.5), h - Inches(1.1), etiqueta,
          tam=13, color=TINTA_2)


def vinetas(s, x, y, w, h, items, tam=16, color_punto=NARANJA):
    pars = [[("■  ", {"color": color_punto, "tam": tam - 4}), (t, {})] if isinstance(t, str) else
            [("■  ", {"color": color_punto, "tam": tam - 4})] + t for t in items]
    caja = texto(s, x, y, w, h, pars, tam=tam, color=TINTA, interlineado=1.15)
    for p in caja.text_frame.paragraphs:
        p.space_after = Pt(10)
    return caja


# --------------------------------------------------------------------------- datos

minam = json.loads((VAL / "minam" / "resumen.json").read_text(encoding="utf-8"))
v = minam["versiones"]
anterior = next(u for k, u in ((u["variante"], u) for u in minam["umbrales"]) if "umbral anterior" in k)
nubes = json.loads((VAL / "comparacion_mascara" / "resumen.json").read_text(encoding="utf-8"))
ultima = almacen.ultima_corrida()
corrida = json.loads((Path(ultima["carpeta"]) / "resumen.json").read_text(encoding="utf-8"))
alertas = corrida["alertas"]
pi = next((a for a in alertas if a["provincia"] == "Puerto Inca" and a["tipo"] == "deforestación"), None)
n_alta = sum(n == "alta" for n in corrida["niveles"].values())
fecha_corrida = corrida["generado"][:10].split("-")[::-1]

# Recortes de las capturas del tablero (capturas a 2x; ver scripts en el manual)
pan = CAP / "panorama.png"
rec_tarjetas = recortar(pan, (725, 480, 2620, 1690), "tarjetas.png")
rec_alertas = recortar(pan, (725, 1720, 2250, 3220), "alertas.png")
rec_mapa = recortar(pan, (750, 3410, 2850, 4390), "mapa.png")
rec_defo = recortar(CAP / "deforestacion.png", (750, 870, 2850, 2400), "deforestacion_mapa.png")
rec_zonas = recortar(CAP / "deforestacion.png", (750, 2440, 2850, 3170), "deforestacion_tabla.png")
cel = Image.open(CAP / "celular.png")
rec_cel = recortar(CAP / "celular.png", (0, 120, cel.width, min(cel.height, 120 + 1500)), "celular.png")
lam = VAL / "minam" / "revision" / "revision_confirmada_1.png"
rec_lam = recortar(lam, (0, 55, 920, 552), "lamina_2zonas.png")

# --------------------------------------------------------------------------- 1. portada

s = prs.slides.add_slide(VACIA)
numero["n"] += 1
rect(s, 0, 0, ANCHO, ALTO, VERDE_OSC)
# Recorte del mapa con la proporción exacta del panel derecho (sin deformar)
mapa = Image.open(rec_mapa)
panel_w, panel_h = ANCHO - Inches(7.1), ALTO
alto_px = mapa.height
ancho_px = int(alto_px * panel_w / panel_h)
x0 = int(mapa.width * 0.42) - ancho_px // 2
portada = TMP / "portada.png"
mapa.crop((max(x0, 0), 0, max(x0, 0) + ancho_px, alto_px)).save(portada)
imagen(s, portada, Inches(7.1), 0, panel_w, panel_h, borde=False)
velo = rect(s, Inches(7.1), 0, panel_w, panel_h, VERDE_OSC)
# python-pptx no expone la transparencia del relleno: se agrega <a:alpha> en el XML (35 % de opacidad)
from pptx.oxml.ns import qn  # noqa: E402
srgb = velo._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
srgb.append(srgb.makeelement(qn("a:alpha"), {"val": "35000"}))
rect(s, Inches(0.7), Inches(0.9), Inches(0.08), Inches(0.5), NARANJA)
texto(s, Inches(0.95), Inches(0.9), Inches(6), Inches(0.5), "INNOBEST HUÁNUCO 2026", tam=16, color=NARANJA,
      negrita=True, anclaje=MSO_ANCHOR.MIDDLE)
texto(s, Inches(0.7), Inches(1.75), Inches(6.2), Inches(3.3), TITULO, tam=30, color=BLANCO, negrita=True,
      interlineado=1.05)
texto(s, Inches(0.7), Inches(4.75), Inches(6.2), Inches(0.5),
      "Incendios · deforestación · salud de la vegetación · contaminación del aire", tam=14,
      color=RGBColor(0xC9, 0xE4, 0xD6))
rect(s, Inches(0.7), Inches(5.45), Inches(5.9), Emu(9525), RGBColor(0x3E, 0x6B, 0x5B))
texto(s, Inches(0.7), Inches(5.6), Inches(6.2), Inches(1.2),
      [[(n, {"negrita": True, "color": BLANCO, "tam": 15})] for n, _ in EQUIPO]
      + [[(EQUIPO[0][1], {"color": RGBColor(0xC9, 0xE4, 0xD6), "tam": 12})]], interlineado=1.2)
s.notes_slide.notes_text_frame.text = (
    "Buenos días. Somos Williams Roberts Jara Daza y Limber Yinier Valdivia Cuzque, de Ingeniería de software con "
    "inteligencia artificial. Presentamos un sistema que vigila con satélites las 11 provincias de Huánuco y avisa "
    "a tiempo sobre incendios, deforestación, deterioro de la vegetación y contaminación del aire.")

# --------------------------------------------------------------------------- 2. problema

s = diapositiva("Las amenazas se detectan tarde y la información está dispersa", "El problema",
                "Huánuco tiene selva, ceja de selva y sierra. Cada semana hay focos de calor, claros nuevos en el "
                "bosque y episodios de humo, pero la información está repartida en varias plataformas técnicas y "
                "se revisa tarde. Estas cifras las midió nuestro propio sistema en la última semana.")
w = Inches(3.85)
tarjeta_dato(s, Inches(0.6), Inches(1.9), w, Inches(2.2), f"{corrida['focos']}",
             f"focos de calor detectados en Huánuco en 7 días (NASA FIRMS)", NARANJA)
tarjeta_dato(s, Inches(0.6) + w + Inches(0.29), Inches(1.9), w, Inches(2.2),
             f"{pi['valor']:,.0f} ha".replace(",", " ") if pi else "—",
             "de posible pérdida de bosque en Puerto Inca en 90 días, frente al mismo periodo de 2025", VERDE)
tarjeta_dato(s, Inches(0.6) + 2 * (w + Inches(0.29)), Inches(1.9), w, Inches(2.2), f"{n_alta} de 11",
             "provincias en alerta alta en la corrida de esta semana", RGBColor(0xD0, 0x3B, 0x3B))
vinetas(s, Inches(0.6), Inches(4.45), Inches(12), Inches(2.4), [
    "Los datos satelitales existen y son gratuitos, pero están en plataformas separadas y requieren conocimiento técnico.",
    "Las nubes de la selva generan falsas alarmas si no se filtran bien: una alerta poco confiable no sirve para actuar.",
    "Las autoridades necesitan saber dónde, cuándo y con qué certeza, no solo un mapa.",
], tam=17)

# --------------------------------------------------------------------------- 3. propuesta

s = diapositiva("Un sistema, cuatro amenazas, once provincias", "Nuestra propuesta",
                "Integramos cuatro fuentes satelitales gratuitas en un solo sistema. Cada amenaza tiene su método "
                "y todas terminan en alertas por provincia con su evidencia. El sistema se actualiza solo cada día.")
amenazas = [
    ("Incendios", "NASA FIRMS\nVIIRS 375 m · MODIS", "Focos de calor agrupados en eventos con severidad", NARANJA),
    ("Deforestación", "Sentinel-2\n10–20 m", "Tala, quema y desbroce dentro del bosque oficial de MINAM", VERDE),
    ("Vegetación", "Sentinel-2\nNDVI mensual", "Deterioro frente al mismo mes del año anterior", RGBColor(0x1B, 0xAF, 0x7A)),
    ("Contaminación", "Sentinel-5P\nNO₂ · CO", "Anomalías semanales de gases por provincia", AZUL),
]
w = Inches(2.88)
for i, (nombre, sat, que, col) in enumerate(amenazas):
    x = Inches(0.6) + i * (w + Inches(0.2))
    rect(s, x, Inches(1.95), w, Inches(4.2), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    rect(s, x, Inches(1.95), w, Inches(0.12), col)
    texto(s, x + Inches(0.25), Inches(2.3), w - Inches(0.5), Inches(0.5), nombre, tam=21, color=VERDE_OSC, negrita=True)
    texto(s, x + Inches(0.25), Inches(2.95), w - Inches(0.5), Inches(0.9), sat, tam=13, color=col, negrita=True)
    texto(s, x + Inches(0.25), Inches(3.95), w - Inches(0.5), Inches(1.9), que, tam=15, color=TINTA_2)
texto(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.5),
      [[("Resultado: ", {"negrita": True, "color": VERDE_OSC}),
        ("alertas por provincia con nivel (alta, media, baja), ubicación, fecha, fuente y calidad de la observación.", {})]],
      tam=15, color=TINTA_2)

# --------------------------------------------------------------------------- 4. arquitectura

s = diapositiva("Cómo funciona", "Arquitectura",
                "El flujo tiene cuatro pasos. Uno: descargamos datos de tres satélites y de los servicios oficiales "
                "de MINAM. Dos: los procesamos con métodos de inteligencia artificial y análisis espacial. Tres: "
                "generamos alertas con evidencia. Cuatro: las mostramos en un tablero web que funciona en celular, "
                "y guardamos cada corrida en una base de datos con respaldo verificado.")
pasos = [
    ("1  Captura", ["NASA FIRMS", "Sentinel-2", "Sentinel-5P", "MINAM Geobosques"]),
    ("2  Procesamiento", ["Nubes: modelo IA", "Mediana por píxel", "Grupos DBSCAN", "Anomalías (z)"]),
    ("3  Alertas", ["Reglas por provincia", "Alta · media · baja", "Evidencia verificable", "Contraste MINAM"]),
    ("4  Uso", ["Tablero web y celular", "Actualización diaria", "Base de datos SQLite", "Respaldo verificado"]),
]
w, h = Inches(2.7), Inches(3.9)
for i, (titulo_paso, items) in enumerate(pasos):
    x = Inches(0.6) + i * (w + Inches(0.44))
    col = [AZUL, VERDE, NARANJA, VERDE_OSC][i]
    rect(s, x, Inches(2.0), w, Inches(0.75), col, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.15)
    texto(s, x, Inches(2.0), w, Inches(0.75), titulo_paso, tam=18, color=BLANCO, negrita=True,
          alinear=PP_ALIGN.CENTER, anclaje=MSO_ANCHOR.MIDDLE)
    rect(s, x, Inches(2.9), w, h - Inches(0.9), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    texto(s, x + Inches(0.25), Inches(3.15), w - Inches(0.5), h - Inches(1.3),
          [[("•  ", {"color": col}), (t, {})] for t in items], tam=15, color=TINTA, interlineado=1.7)
    if i < 3:
        flecha = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + w + Inches(0.08), Inches(2.2), Inches(0.28),
                                    Inches(0.35))
        flecha.fill.solid(); flecha.fill.fore_color.rgb = APAGADO; flecha.line.fill.background()
texto(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.5),
      "Python · Streamlit · Copernicus Data Space · NASA FIRMS · servicios abiertos de MINAM · sin servicios de pago",
      tam=13, color=APAGADO, alinear=PP_ALIGN.CENTER)

# --------------------------------------------------------------------------- 5. inteligencia

s = diapositiva("¿Dónde está la inteligencia del sistema?", "Inteligencia artificial y análisis de datos",
                "La inteligencia no es un solo modelo: es una cadena de técnicas. La detección de nubes usa un "
                "modelo de aprendizaje automático, s2cloudless. Los focos se agrupan con DBSCAN, un algoritmo de "
                "agrupamiento no supervisado. Comparamos cada píxel con el mismo periodo del año anterior y "
                "detectamos anomalías estadísticas en la contaminación. Y todo se contrasta con datos oficiales.")
tecnicas = [
    ("Aprendizaje automático", "s2cloudless clasifica nubes y sus bordes en cada imagen: eliminó el 100 % de las "
                               "falsas alarmas por nubes de la muestra."),
    ("Agrupamiento no supervisado", "DBSCAN con distancia geográfica une focos de calor cercanos en eventos de "
                                    "incendio y estima su severidad."),
    ("Análisis multitemporal", "Mediana por píxel de todas las pasadas sin nubes, comparada con el mismo periodo "
                               "del año anterior para no confundir la estación seca."),
    ("Detección de anomalías", "Puntaje z de NO₂ y CO frente a la línea base de cada provincia."),
    ("Fusión con datos oficiales", "Máscara de bosque y alertas de MINAM para filtrar cultivos y validar."),
    ("Calibración basada en evidencia", "Cada umbral se eligió midiendo precisión y sensibilidad, no a ojo."),
]
for i, (t, d) in enumerate(tecnicas):
    col, fila = i % 2, i // 2
    x = Inches(0.6) + col * Inches(6.15)
    y = Inches(1.9) + fila * Inches(1.6)
    rect(s, x, y, Inches(5.95), Inches(1.42), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    rect(s, x, y + Inches(0.25), Inches(0.07), Inches(0.9), [VERDE, NARANJA, AZUL][fila])
    texto(s, x + Inches(0.3), y + Inches(0.18), Inches(5.4), Inches(0.4), t, tam=16, color=VERDE_OSC, negrita=True)
    texto(s, x + Inches(0.3), y + Inches(0.58), Inches(5.45), Inches(0.8), d, tam=13, color=TINTA_2)

# --------------------------------------------------------------------------- 6. tablero

s = diapositiva("El tablero: lo importante primero", "Demostración",
                "Al abrir el tablero se ve en segundos el estado de cada provincia, ordenadas de mayor a menor "
                "alerta, y las alertas prioritarias con su fecha, fuente y calidad. Abre con la última corrida "
                "automática, así que no depende de la conexión en el momento. Aquí conviene pasar a la demostración "
                "en vivo.")
imagen_ajustada(s, rec_tarjetas, Inches(0.6), Inches(1.75), Inches(6.0), Inches(5.0))
imagen_ajustada(s, rec_alertas, Inches(6.85), Inches(1.75), Inches(5.9), Inches(5.0))

# --------------------------------------------------------------------------- 7. deforestación

s = diapositiva("Deforestación: dónde, cuánto y con qué respaldo", "Demostración",
                "En la sección de deforestación elegimos una provincia; en Puerto Inca el sistema marca las zonas de "
                "posible pérdida dentro del bosque oficial. Cada zona trae coordenadas, hectáreas y si coincide con "
                "una alerta oficial de MINAM, para que un guardaparque pueda ir directamente al lugar.")
imagen_ajustada(s, rec_defo, Inches(0.6), Inches(1.75), Inches(7.3), Inches(5.05))
rect(s, Inches(8.2), Inches(1.75), Inches(4.53), Inches(5.05), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.04)
texto(s, Inches(8.45), Inches(1.95), Inches(4.1), Inches(0.4), "Cada zona incluye", tam=17, color=VERDE_OSC,
      negrita=True)
vinetas(s, Inches(8.45), Inches(2.5), Inches(4.1), Inches(4.2), [
    "Coordenadas y hectáreas",
    "Coincidencia con alerta oficial MINAM (sí / no)",
    "Imágenes de antes y después sin nubes",
    "Porcentaje de la provincia observado sin nubes",
    "Exportación a CSV para el trabajo de campo",
], tam=15)

# --------------------------------------------------------------------------- 8. validación MINAM

s = diapositiva("Validación independiente con datos oficiales del Estado", "Resultados",
                f"Esta es la cifra principal. Comparamos nuestras alertas con las alertas tempranas oficiales de "
                f"MINAM, que usan otro satélite y otro método. El {v['v3']['precision_zonas_pct']:.0f} % de nuestras "
                f"zonas coincide con una alerta oficial; si el sistema marcara zonas al azar, coincidiría solo el "
                f"{v['v3']['azar_pixeles_pct']:.0f} %. Cualquiera puede repetir esta comparación en Geobosques.")
rect(s, Inches(0.6), Inches(1.85), Inches(3.9), Inches(4.9), VERDE_OSC, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
texto(s, Inches(0.85), Inches(2.15), Inches(3.4), Inches(1.4), f"{v['v3']['precision_zonas_pct']:.0f} %", tam=72,
      color=BLANCO, negrita=True)
texto(s, Inches(0.85), Inches(3.6), Inches(3.4), Inches(1.4),
      "de las zonas de pérdida detectadas coinciden con alertas oficiales de MINAM", tam=16,
      color=RGBColor(0xC9, 0xE4, 0xD6))
texto(s, Inches(0.85), Inches(5.15), Inches(3.4), Inches(1.4),
      [[("Por azar: ", {"negrita": True, "color": BLANCO}),
        (f"{v['v3']['azar_pixeles_pct']:.0f} %", {"color": BLANCO})],
       [("Otro satélite (Landsat) y otro método: validación independiente.", {"color": RGBColor(0xC9, 0xE4, 0xD6),
                                                                              "tam": 13})]], tam=15)
imagen_ajustada(s, INN / "fig_precision_versiones.png", Inches(4.8), Inches(1.85), Inches(7.93), Inches(4.9))

# --------------------------------------------------------------------------- 9. mejora medible

s = diapositiva("Cada mejora se midió antes de adoptarla", "Resultados",
                "No agregamos más inteligencia artificial por agregar. Primero corregimos lo que fallaba: la máscara "
                "de nubes eliminó las falsas alarmas por nubes, el mapa oficial de bosque eliminó las parcelas "
                "agrícolas y el ajuste del umbral recuperó claros con rebrote. El sistema prioriza no dar falsas "
                "alarmas.")
imagen_ajustada(s, INN / "fig_precision_sensibilidad.png", Inches(0.6), Inches(1.75), Inches(7.6), Inches(5.05))
w = Inches(4.3)
tarjeta_dato(s, Inches(8.43), Inches(1.75), w, Inches(1.55),
             f"{v['v1']['precision_zonas_pct']:.0f} % → {v['v3']['precision_zonas_pct']:.0f} %",
             "precisión frente a MINAM", VERDE_OSC, 34)
tarjeta_dato(s, Inches(8.43), Inches(3.5), w, Inches(1.55),
             f"{nubes['perdida']['zonas'] - nubes['perdida']['confirmadas']} → 0",
             "falsas alarmas en la muestra revisada de la primera versión", NARANJA, 34)
tarjeta_dato(s, Inches(8.43), Inches(5.25), w, Inches(1.55),
             f"{v['v3']['sensibilidad_alertas_grandes_pct']:.0f} %",
             "de los claros grandes registrados por MINAM detectados", AZUL, 34)

# --------------------------------------------------------------------------- 10. evidencia visual

s = diapositiva("Evidencia visual: antes y después, sin nubes", "Resultados",
                "Así se ve la evidencia. A la izquierda el bosque en 2025, luego el mismo lugar en 2026, en falso "
                "color el suelo desnudo resalta en rosado y a la derecha el contorno que detectó el sistema. Revisamos "
                "60 zonas sorteadas al azar con apoyo de inteligencia artificial y no encontramos falsas alarmas; esa "
                "revisión es preliminar y la estamos completando con verificación del equipo.")
imagen_ajustada(s, rec_lam, Inches(0.6), Inches(1.75), Inches(8.2), Inches(5.05))
rect(s, Inches(9.05), Inches(1.75), Inches(3.68), Inches(5.05), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.04)
texto(s, Inches(9.3), Inches(2.0), Inches(3.2), Inches(0.9), "0", tam=60, color=VERDE_OSC, negrita=True)
texto(s, Inches(9.3), Inches(3.05), Inches(3.2), Inches(0.9), "falsas alarmas en 60 zonas sorteadas al azar",
      tam=15, color=TINTA_2)
texto(s, Inches(9.3), Inches(4.35), Inches(3.2), Inches(2.3),
      "Revisión visual asistida por IA, preliminar. La evidencia principal es la comparación con MINAM.",
      tam=12, color=APAGADO)

# --------------------------------------------------------------------------- 11. operación y celular

s = diapositiva("Listo para usar: en la oficina y en el campo", "Operación",
                "El sistema funciona en una computadora común. Se actualiza solo cada día, cada fuente por separado, "
                "de modo que si un servicio falla los demás siguen disponibles. Cada corrida se guarda en una base de "
                "datos con respaldo verificado, y el tablero se adapta al celular para consultarlo desde el campo.")
imagen_ajustada(s, rec_cel, Inches(0.6), Inches(1.7), Inches(3.3), Inches(5.15))
vinetas(s, Inches(4.4), Inches(1.95), Inches(8.3), Inches(4.9), [
    [("Abre en segundos ", {"negrita": True}), ("con la última corrida automática; consulta satélites solo si se pide.", {})],
    [("Actualización diaria automática ", {"negrita": True}), ("y cada fuente independiente: si una falla, se usan sus últimos datos buenos con su fecha.", {})],
    [("Historial en base de datos ", {"negrita": True}), ("para ver la evolución de las alertas.", {})],
    [("Respaldo verificado: ", {"negrita": True}), ("se comprueba que cada copia se puede restaurar.", {})],
    [("Costo cero: ", {"negrita": True}), ("satélites y servicios gratuitos, software libre.", {})],
], tam=17)

# --------------------------------------------------------------------------- 12. impacto

s = diapositiva("Impacto e innovación", "Valor del proyecto",
                "¿Quién lo usa? Autoridades forestales y ambientales, municipalidades y comunidades. ¿Qué lo hace "
                "innovador? Que integra las cuatro amenazas en un solo lugar, que cada alerta es verificable y que su "
                "precisión está medida con datos oficiales, no solo declarada.")
cols = [
    ("Para quién", ["Autoridades forestales (SERFOR, ATFFS)", "Gobierno regional y municipalidades",
                    "Comunidades y guardaparques", "Universidades e investigadores"], VERDE),
    ("Qué lo hace innovador", ["Cuatro amenazas en un solo tablero", "Alertas con evidencia verificable",
                               "Precisión medida con datos oficiales", "Funciona sin servicios de pago"], NARANJA),
    ("Qué permite", ["Llegar antes a un incendio o una tala", "Priorizar patrullajes con coordenadas",
                     "Seguir la evolución semana a semana", "Sustentar denuncias con imágenes"], AZUL),
]
w = Inches(3.9)
for i, (t, items, col) in enumerate(cols):
    x = Inches(0.6) + i * (w + Inches(0.2))
    rect(s, x, Inches(1.9), w, Inches(4.85), BLANCO, LINEA, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    rect(s, x, Inches(1.9), w, Inches(0.12), col)
    texto(s, x + Inches(0.3), Inches(2.25), w - Inches(0.6), Inches(0.5), t, tam=20, color=VERDE_OSC, negrita=True)
    vinetas(s, x + Inches(0.3), Inches(2.95), w - Inches(0.6), Inches(3.7), items, tam=15, color_punto=col)

# --------------------------------------------------------------------------- 13. límites y próximos pasos

s = diapositiva("Límites conocidos y próximos pasos", "Hoja de ruta",
                "Somos transparentes con los límites. Las nubes de la selva alta impiden observar parte del territorio; "
                "el próximo paso es incorporar radar Sentinel-1, que ve a través de las nubes. También queremos avisos "
                "automáticos por correo o WhatsApp y una verificación de campo con las autoridades.")
texto(s, Inches(0.6), Inches(1.85), Inches(5.9), Inches(0.5), "Límites actuales", tam=19, color=VERDE_OSC, negrita=True)
vinetas(s, Inches(0.6), Inches(2.45), Inches(5.9), Inches(4.4), [
    "Las nubes impiden observar parte de la selva alta en cada periodo (el sistema informa qué porcentaje observó).",
    "No evalúa bosques secundarios ni andinos que no están en el mapa oficial de MINAM.",
    "Prioriza la precisión: detecta mejor los claros grandes que los pequeños.",
    "El aire se mide en la columna atmosférica; no hay estaciones en tierra para contrastar.",
], tam=18, color_punto=APAGADO)
texto(s, Inches(6.85), Inches(1.85), Inches(5.9), Inches(0.5), "Próximos pasos", tam=19, color=VERDE_OSC, negrita=True)
vinetas(s, Inches(6.85), Inches(2.45), Inches(5.9), Inches(4.4), [
    "Radar Sentinel-1 para ver a través de las nubes.",
    "Avisos automáticos por correo o WhatsApp ante alertas altas.",
    "Verificación de campo con autoridades forestales.",
    "Modelos de aprendizaje profundo cuando existan datos verificados para entrenarlos.",
], tam=18, color_punto=VERDE)

# --------------------------------------------------------------------------- 14. cierre

s = prs.slides.add_slide(VACIA)
numero["n"] += 1
rect(s, 0, 0, ANCHO, ALTO, VERDE_OSC)
texto(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.2), "Gracias", tam=54, color=BLANCO, negrita=True)
texto(s, Inches(0.8), Inches(2.65), Inches(11.7), Inches(1.4),
      f"Alertas tempranas verificables para Huánuco: {v['v3']['precision_zonas_pct']:.0f} % de coincidencia con "
      f"las alertas oficiales de MINAM.", tam=22, color=RGBColor(0xC9, 0xE4, 0xD6))
rect(s, Inches(0.8), Inches(4.35), Inches(0.08), Inches(1.35), NARANJA)
texto(s, Inches(1.1), Inches(4.3), Inches(11), Inches(1.5),
      [[(n, {"negrita": True, "color": BLANCO, "tam": 18})] for n, _ in EQUIPO]
      + [[(EQUIPO[0][1], {"color": RGBColor(0xC9, 0xE4, 0xD6), "tam": 14})]], interlineado=1.25)
texto(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5),
      "INNOBEST HUÁNUCO 2026  ·  Fuentes: NASA FIRMS, Copernicus Sentinel-2 y Sentinel-5P, MINAM Geobosques",
      tam=12, color=RGBColor(0x9C, 0xC2, 0xB0))
s.notes_slide.notes_text_frame.text = ("Gracias. Quedamos atentos a sus preguntas; si lo desean, mostramos el "
                                       "tablero en vivo.")

ruta = INN / "Presentacion_INNOBEST_2026.pptx"
prs.save(ruta)
print(ruta, f"({numero['n']} diapositivas)")
